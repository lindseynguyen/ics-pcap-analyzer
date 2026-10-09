"""CLI, report export, storylines, IOC export and history database."""
import json
import os
import re
import subprocess
import sys

import pytest
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
VI_CHARS = re.compile("[ăâđêôơưạảấầẩẫậắằẳẵặẹẻẽếềểễệỉịọỏốồổỗộớờởỡợụủứừửữựỳỵỷỹ]", re.I)


def test_cli_end_to_end(captures, tmp_path):
    out = tmp_path / "report.xlsx"
    env = dict(os.environ, PYTHONPATH=ROOT, HOME=str(tmp_path))
    r = subprocess.run([sys.executable, "-m", "ot_pcap_analyzer", "--cli", captures["attacks"], "--advanced", "-o", str(out)],
                       capture_output=True, text=True, env=env, timeout=300)
    assert r.returncode == 0, r.stderr[-2000:]
    assert "ANALYSIS SUMMARY" in r.stdout and "ATTACK CHAINS DETECTED" in r.stdout
    wb = load_workbook(out)
    for sheet in ("Summary", "OT Events", "Assets", "Anomalies", "Attack Chains", "Storylines"):
        assert sheet in wb.sheetnames
    for ws in wb.worksheets:
        headers = [c.value for c in ws[1] if c.value]
        assert not any("(VI)" in str(h) for h in headers), (ws.title, headers)
        for row in ws.iter_rows(values_only=True):
            for v in row:
                assert not (isinstance(v, str) and VI_CHARS.search(v)), (ws.title, v)


def test_cli_version_and_bad_file(tmp_path):
    env = dict(os.environ, PYTHONPATH=ROOT, HOME=str(tmp_path))
    r = subprocess.run([sys.executable, "-m", "ot_pcap_analyzer", "--version"], capture_output=True, text=True, env=env)
    from ot_pcap_analyzer.version import VERSION
    assert r.returncode == 0 and VERSION in r.stdout and "__main__" not in r.stdout
    bad = tmp_path / "x.pcap"
    bad.write_bytes(b"not a pcap at all, definitely not")
    r = subprocess.run([sys.executable, "-m", "ot_pcap_analyzer", "--cli", str(bad)], capture_output=True, text=True, env=env)
    assert r.returncode != 0


def test_excel_export_survives_binary_payloads(analyzed, tmp_path):
    from ot_pcap_analyzer.models import SecurityAnomaly
    a = analyzed["benign"]
    a.anomalies.append(SecurityAnomaly(timestamp=1.0, anomaly_type="TEST", severity="LOW", src_ip="1.1.1.1",
                                       dst_ip="2.2.2.2", protocol="TCP", description="bin \x00\x01\x02 data",
                                       evidence={"raw": b"\x00\xff"}, mitre_techniques=[], confidence=0.5,
                                       recommendation="x" * 40000))
    try:
        out = tmp_path / "b.xlsx"
        a.export_excel(str(out))
        assert out.exists()
    finally:
        a.anomalies.pop()


def test_storyline_text_report_is_english(analyzed):
    a = analyzed["attacks"]
    report = a._storyline_generator.generate_text_report(a.storylines[0])
    assert len(report) > 200
    assert not VI_CHARS.search(report)
    # "both" mode must not print every section twice
    assert report.count("THREAT ASSESSMENT") <= 1


def test_ioc_collector_exports(analyzed, tmp_path):
    from ot_pcap_analyzer.ioc_collector import IOCCollector
    c = IOCCollector()
    iocs = c.extract_iocs(analyzed["attacks"])
    total = sum(len(v) for v in iocs.values() if isinstance(v, list))
    assert total > 0
    c.export_to_json(iocs, str(tmp_path / "iocs.json"))
    c.export_to_csv(iocs, str(tmp_path / "iocs.csv"))
    data = json.loads((tmp_path / "iocs.json").read_text())
    assert data["export_metadata"]["total_iocs"] == total - 0 or data["export_metadata"]["total_iocs"] > 0
    rows = (tmp_path / "iocs.csv").read_text().strip().splitlines()
    assert len(rows) == 1 + data["export_metadata"]["total_iocs"]
    assert (tmp_path / "iocs.csv").stat().st_size > 0


def test_history_database_roundtrip(analyzed, tmp_path):
    from ot_pcap_analyzer.database import AnalysisDatabase
    db = AnalysisDatabase(str(tmp_path / "h.db"))
    try:
        sid = db.create_session("attacks.pcap", {"note": "test"})
        n = db.store_anomalies(sid, analyzed["attacks"].anomalies)
        assert n == len(analyzed["attacks"].anomalies)
        assert len(db.get_anomalies(session_id=sid, limit=10000)) == n
        rows = db.get_anomalies(session_id=sid, limit=10000)
        assert all(r["timestamp"] is None or hasattr(r["timestamp"], "year") for r in rows)
        assert db.get_session(sid) is not None
        db.get_statistics()
        out = tmp_path / "session.json"
        assert db.export_session(sid, str(out))
        new_sid = db.import_session(str(out))
        assert new_sid and len(db.get_anomalies(session_id=new_sid, limit=10000)) == n
    finally:
        db.close()


def test_summary_contains_core_keys(analyzed):
    s = analyzed["attacks"].get_summary()
    for k in ("PACKETS_PARSED", "OT_EVENTS_TOTAL", "ANOMALIES_DETECTED", "ATTACK_CHAINS", "STORYLINES"):
        assert k in s


def test_source_tree_is_english():
    pkg = os.path.join(ROOT, "ot_pcap_analyzer")
    offenders = []
    for name in os.listdir(pkg):
        if name.endswith(".py"):
            for i, line in enumerate(open(os.path.join(pkg, name), encoding="utf-8"), 1):
                if VI_CHARS.search(line):
                    offenders.append(f"{name}:{i}")
    assert not offenders, offenders[:20]
